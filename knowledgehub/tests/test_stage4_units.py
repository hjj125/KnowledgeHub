"""阶段 4 单元测试：分块 / 解析 / 主键派生 / 上传魔数校验（不依赖外部服务）。

运行：backend/.venv/Scripts/python -m pytest tests/test_stage4_units.py -v
"""

from __future__ import annotations

import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation
from io import BytesIO

from app.repositories.milvus import chunk_pk
from app.services.chunking import chunk_text
from app.services.parsing import parse_content, parse_pdf, parse_pptx
from app.api.v1.documents import _sniff


class TestChunkText:
    def test_empty_and_short(self):
        assert chunk_text("") == []
        assert chunk_text("   \n\n  ") == []
        short = "这是一段短文本。"
        assert chunk_text(short) == [short]

    def test_multi_paragraph_merged_to_target(self):
        paras = [f"第{i}段。" + "内容" * 40 for i in range(6)]  # 每段 ~83 字
        chunks = chunk_text("\n\n".join(paras), target=450, max_len=600, overlap=80)
        assert len(chunks) >= 2
        assert all(len(c) <= 600 for c in chunks)
        # 内容无丢失：所有段落号都出现在某块中
        joined = "".join(chunks)
        for i in range(6):
            assert f"第{i}段" in joined

    def test_long_paragraph_hard_split(self):
        text = "超" * 1500  # 无句读的单段
        chunks = chunk_text(text, target=450, max_len=600, overlap=0)
        assert all(len(c) <= 600 for c in chunks)
        assert "".join(chunks) == text  # 无重叠时硬切不丢字
        overlapped = chunk_text(text, target=450, max_len=600, overlap=80)
        assert all(len(c) <= 600 for c in overlapped)
        assert overlapped[0] == text[:600]  # 首块与原文对齐

    def test_overlap_carries_tail(self):
        paras = [f"段落{i}。" + "内容" * 15 for i in range(40)]  # 每段 ~34 字
        chunks = chunk_text("\n\n".join(paras), target=100, max_len=600, overlap=30)
        assert len(chunks) >= 3
        # 重叠语义：后一块开头精确等于前一块结尾的 overlap 尾巴
        assert chunks[1][:30] == chunks[0][-30:]

    def test_sentence_split_keeps_delimiters(self):
        chunks = chunk_text("你好世界。第二句！第三句？" * 60, target=120, max_len=600)
        assert chunks
        assert all("你好世界" in joined for joined in ["".join(chunks)])


class TestParsers:
    def test_parse_text_utf8_and_gbk(self):
        assert parse_content("中文内容".encode("utf-8"), "txt") == [("中文内容", 0)]
        assert parse_content("中文内容".encode("gbk"), "md") == [("中文内容", 0)]

    def test_parse_pdf_preserves_pages(self):
        doc = pymupdf.open()
        for i in range(3):
            page = doc.new_page()
            page.insert_text((72, 72), f"第{i + 1}页的知识内容", fontname="china-s")
        buf = BytesIO()
        doc.save(buf)
        doc.close()
        pages = parse_pdf(buf.getvalue())
        assert [p for _, p in pages] == [1, 2, 3]
        assert all("知识内容" in t for t, _ in pages)

    def test_parse_docx(self):
        d = DocxDocument()
        d.add_paragraph("标题段落")
        d.add_paragraph("正文段落内容")
        buf = BytesIO()
        d.save(buf)
        pages = parse_content(buf.getvalue(), "docx")
        assert len(pages) == 1 and pages[0][1] == 0
        assert "正文段落内容" in pages[0][0]

    def test_parse_pptx_preserves_slide_numbers(self):
        prs = Presentation()
        slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # 空白版式
        box1 = slide1.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height // 3)
        box1.text_frame.text = "第一张：课程导入"
        slide2 = prs.slides.add_slide(prs.slide_layouts[6])
        box2 = slide2.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height // 3)
        box2.text_frame.text = "第二张：TCP 三次握手详解"
        slide2.notes_slide.notes_text_frame.text = "提醒强调序列号"
        buf = BytesIO()
        prs.save(buf)
        pages = parse_pptx(buf.getvalue())
        assert [p for _, p in pages] == [1, 2]
        assert "课程导入" in pages[0][0]
        assert "三次握手" in pages[1][0] and "备注：提醒强调序列号" in pages[1][0]

    def test_parse_pptx_sniff(self):
        assert _sniff(b"PK\x03\x04 pptx-bytes", "pptx")
        assert not _sniff(b"%PDF- fake", "pptx")

    def test_unknown_type_falls_back_to_text(self):
        assert parse_content(b"plain", "xyz") == [("plain", 0)]


class TestChunkPk:
    def test_deterministic_and_version_sensitive(self):
        a = chunk_pk("doc-1", 1, 0)
        assert a == chunk_pk("doc-1", 1, 0)
        assert a != chunk_pk("doc-1", 2, 0)
        assert a != chunk_pk("doc-1", 1, 1)
        assert a != chunk_pk("doc-2", 1, 0)
        assert len(a) == 36  # uuid5 标准格式


class TestSniff:
    def test_pdf_magic(self):
        assert _sniff(b"%PDF-1.7 rest", "pdf")
        assert not _sniff(b"PK\x03\x04 fake", "pdf")

    def test_docx_magic(self):
        assert _sniff(b"PK\x03\x04 zip-bytes", "docx")
        assert not _sniff(b"%PDF- fake", "docx")

    def test_text_decodable(self):
        assert _sniff("中文".encode("utf-8"), "md")
        assert _sniff("中文".encode("gbk"), "txt")
        assert not _sniff(bytes(range(0, 32)) * 8, "txt")

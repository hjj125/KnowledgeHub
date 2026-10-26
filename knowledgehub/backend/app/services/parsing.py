"""文件解析（§5.3）：PDF/DOCX/PPTX/MD/TXT → [(文本, 页码)]。

PDF 逐页提取并保留页码（来源定位依赖它）；DOCX 无固定分页概念记 0；
PPTX 以幻灯片编号作页码（检索可定位到第 N 张）；文本类整体作为一段。
图片原件解析列入 P2。
"""

from io import BytesIO

import pymupdf
from docx import Document as DocxDocument
from pptx import Presentation


def parse_pdf(content: bytes) -> list[tuple[str, int]]:
    doc = pymupdf.open(stream=content, filetype="pdf")
    try:
        pages = []
        for page_no, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append((text, page_no))
        return pages
    finally:
        doc.close()


def parse_docx(content: bytes) -> list[tuple[str, int]]:
    document = DocxDocument(BytesIO(content))
    text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
    return [(text, 0)] if text.strip() else []


def _shape_text(shape) -> list[str]:
    """递归收集形状文本：文本框 / 组合形状 / 表格。"""
    lines: list[str] = []
    if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):  # GROUP
        for sub in shape.shapes:
            lines.extend(_shape_text(sub))
        return lines
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs).strip()
            if text:
                lines.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                lines.append(line)
    return lines


def parse_pptx(content: bytes) -> list[tuple[str, int]]:
    """逐张幻灯片提取：文本框/组合/表格 + 演讲者备注；幻灯片编号即页码。"""
    prs = Presentation(BytesIO(content))
    pages: list[tuple[str, int]] = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_text(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"备注：{notes}")
        text = "\n".join(lines).strip()
        if text:
            pages.append((text, slide_no))
    return pages


def parse_text(content: bytes) -> list[tuple[str, int]]:
    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("gbk", errors="replace")
    return [(text, 0)] if text.strip() else []


def parse_content(content: bytes, file_type: str) -> list[tuple[str, int]]:
    if file_type == "pdf":
        return parse_pdf(content)
    if file_type == "docx":
        return parse_docx(content)
    if file_type == "pptx":
        return parse_pptx(content)
    return parse_text(content)

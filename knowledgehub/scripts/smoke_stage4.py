"""阶段 4 冒烟测试：文档上传与语义检索全链路。

前置：后端已启动、MongoDB/Redis/Neo4j/Milvus healthy；docker CLI 可用（停机降级场景）。
本脚本自行拉起/终止 publisher 与两个 Worker 子进程，用 pymilvus 同步客户端直查向量断言。

场景（§5.3 / §5.6 / §9.5 / §10.5）：
  1. 笔记入向量库 → 语义检索命中（snippet 来自 MongoDB chunks 回填）
  2. 笔记更新 → 旧版本命中被版本护栏剔除（新内容可检索）
  3. 元数据模式检索（无向量依赖）
  4. 多用户隔离：B 用户检索不到 A 用户的向量
  5. TXT/PDF 上传 → 解析分块 → 嵌入 → ready，检索命中且带页码
  6. 上传校验：扩展名白名单 / 超限 413；空白文档 ready(空)
  7. 同字节文件 GridFS 去重（fs.files 单份、documents 两份）
  8. reindex → content_version+1 → 重新索引后仍可检索
  9. 删笔记/删文档 → Milvus 向量与 Mongo chunks 清空（GridFS 引用计数）
  10. Milvus 停机 → /search 自动降级 metadata（degraded=true），恢复后正常
"""

from __future__ import annotations

import io
import os
import subprocess
import sys
import time
import uuid
from pathlib import Path

import httpx
import pymupdf
import redis as redis_sync
from pymongo import MongoClient

PROJECT = Path(__file__).resolve().parents[1]
BACKEND = PROJECT / "backend"
STREAM = "stream:knowledge_events"
SUBPROCESS_ENV = {**os.environ, "PYTHONPATH": f"{PROJECT}{os.pathsep}{BACKEND}"}

PASSED: list[str] = []
FAILED: list[str] = []


def check(name: str, condition: bool, detail: str = "") -> None:
    if condition:
        PASSED.append(name)
        print(f"  √ {name}" + (f"  [{detail}]" if detail else ""))
    else:
        FAILED.append(name)
        print(f"  × {name}  [{detail}]")


def wait_until(desc: str, fn, timeout: float = 60.0, interval: float = 0.5) -> bool:
    deadline = time.time() + timeout
    while time.time() < deadline:
        try:
            if fn():
                return True
        except Exception:  # noqa: BLE001 — 同步期间抖动属预期
            pass
        time.sleep(interval)
    print(f"    (timeout waiting: {desc})")
    return False


def start_worker(module: str, tag: str) -> subprocess.Popen:
    logs = PROJECT / "tmp"
    logs.mkdir(exist_ok=True)
    log_file = open(logs / f"smoke4_{tag}.log", "w", encoding="utf-8")  # noqa: SIM115
    return subprocess.Popen(
        [sys.executable, "-m", module],
        cwd=str(BACKEND),
        env=SUBPROCESS_ENV,
        stdout=log_file,
        stderr=subprocess.STDOUT,
    )


def stop_proc(proc: subprocess.Popen) -> None:
    if proc.poll() is None:
        proc.terminate()
        try:
            proc.wait(timeout=5)
        except subprocess.TimeoutExpired:
            proc.kill()


def make_pdf(pages_text: list[str]) -> bytes:
    doc = pymupdf.open()
    for text in pages_text:
        page = doc.new_page()
        page.insert_text((72, 90), text, fontname="china-s", fontsize=14)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def make_pptx(slides_text: list[str]) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for text in slides_text:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(0, 0, prs.slide_width, prs.slide_height // 3)
        box.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def main() -> int:
    base = "http://127.0.0.1:8000"
    api = f"{base}/api/v1"
    http = httpx.Client(timeout=30)
    mongo = MongoClient("mongodb://localhost:27017/?directConnection=true", tz_aware=True)
    db = mongo["knowledgehub"]
    rds = redis_sync.Redis.from_url("redis://localhost:6379/0", decode_responses=True, socket_timeout=5)

    from pymilvus import MilvusClient

    mv = MilvusClient(uri="http://localhost:19530")
    COLLECTION = "knowledge_chunks"

    def mv_rows(doc_id: str) -> list[dict]:
        return mv.query(collection_name=COLLECTION, filter=f'document_id == "{doc_id}"',
                        output_fields=["pk", "content_version", "page"])

    def search(token_headers: dict, body: dict) -> dict:
        r = http.post(f"{api}/search", headers=token_headers, json=body)
        assert r.status_code == 200, r.text
        return r.json()

    def top_result(token_headers: dict, query: str, **kw) -> dict | None:
        body = {"query": query, "top_k": 5, **kw}
        return search(token_headers, body)["results"][0] if search(token_headers, body)["results"] else None

    def doc_status(doc_id: str, token_headers: dict) -> dict:
        return http.get(f"{api}/documents/{doc_id}", headers=token_headers).json()

    def wait_ready(doc_id: str, token_headers: dict, min_indexed_version: int = 1) -> bool:
        return wait_until(
            f"document {doc_id[:8]} ready",
            lambda: doc_status(doc_id, token_headers).get("status") == "ready"
            and doc_status(doc_id, token_headers).get("indexed_version", 0) >= min_indexed_version,
        )

    print(f"== 阶段 4 冒烟测试 ({api}) ==")

    # --- 准备：两个用户、知识库，启动 publisher + workers ---
    suffix = uuid.uuid4().hex[:8]
    password = "smoke-passw0rd"

    def register(tag: str) -> dict:
        email = f"smoke4_{tag}_{suffix}@test.dev"
        http.post(f"{api}/auth/register", json={"email": email, "password": password})
        token = http.post(f"{api}/auth/login", json={"email": email, "password": password}).json()["access_token"]
        return {"Authorization": f"Bearer {token}"}

    auth = register("a")
    auth_b = register("b")
    kb = http.post(f"{api}/knowledge-bases", headers=auth, json={"name": "计算机科学"}).json()

    publisher = start_worker("workers.event_publisher", "publisher")
    graph = start_worker("workers.graph_indexer", "graph")
    vector = start_worker("workers.vector_indexer", "vector")

    try:
        # --- 1. 笔记 → 向量 → 语义检索 ---
        note_a = http.post(f"{api}/notes", headers=auth, json={
            "kb_id": kb["id"], "title": "进程与线程",
            "content_md": "进程是资源分配的基本单位。线程是CPU调度的基本单位，同一进程内的线程共享地址空间。",
        }).json()
        note_b = http.post(f"{api}/notes", headers=auth, json={
            "kb_id": kb["id"], "title": "死锁",
            "content_md": "死锁的四个必要条件：互斥、持有并等待、不可剥夺、循环等待。",
        }).json()
        ok = wait_until("笔记A向量就位", lambda: len(mv_rows(note_a["id"])) >= 1)
        check("笔记创建 → 分块嵌入 → Milvus 向量就位", ok, f"n={len(mv_rows(note_a['id']))}")
        check("Mongo chunks 权威分块落库", db["chunks"].count_documents({"document_id": note_a["id"]}) >= 1)

        hit = top_result(auth, "线程是什么概念")
        check("语义检索命中笔记A（note 源，含 snippet）",
              hit is not None and hit["source_id"] == note_a["id"] and hit["source_type"] == "note"
              and "CPU调度" in hit["snippet"], hit and f"score={hit['score']:.3f}" if hit else "no-hit")

        hit_b = top_result(auth, "互斥和循环等待")
        check("语义检索命中笔记B", hit_b is not None and hit_b["source_id"] == note_b["id"],
              hit_b and f"score={hit_b['score']:.3f}" if hit_b else "no-hit")

        # --- 2. 版本护栏：更新后旧命中剔除 ---
        http.patch(f"{api}/notes/{note_a['id']}", headers=auth, json={
            "title": "进程与线程",
            "content_md": "进程切换涉及上下文保存与恢复，页表切换开销显著。", "expected_content_version": 1,
        })
        ok = wait_until("新内容可检索", lambda: (lambda h: h is not None and "上下文保存" in h["snippet"])(
            top_result(auth, "进程切换的开销")))
        check("笔记更新 → 新内容可语义检索", ok)

        stale_results = search(auth, {"query": "线程是什么概念", "top_k": 10})["results"]
        check("旧版本向量被替换（版本护栏，无陈旧命中）",
              all(r["source_id"] != note_a["id"] or "CPU调度" not in r["snippet"] for r in stale_results),
              f"n={len(stale_results)}")

        # --- 3. 元数据模式 ---
        r = search(auth, {"query": "死锁", "mode": "metadata"})
        check("元数据模式：关键词命中笔记B（degraded=false）",
              r["mode"] == "metadata" and not r["degraded"]
              and any(x["source_id"] == note_b["id"] for x in r["results"]))

        # --- 4. 多用户隔离 ---
        r_b = search(auth_b, {"query": "线程是什么概念", "top_k": 10})
        check("多用户隔离：B 用户检索不到 A 用户向量", r_b["results"] == [])

        # --- 5. 文档上传 → ready → 带页码检索 ---
        pdf_bytes = make_pdf([
            "KnowledgeHub 测试文档第一页：计算机网络概述。",
            "第二页内容：TCP 三次握手建立连接的详细过程与序列号确认机制。",
            "第三页内容：HTTP 协议演进与 TLS 加密握手。",
        ])
        r = http.post(f"{api}/knowledge-bases/{kb['id']}/documents", headers=auth,
                      files={"file": ("计算机网络.pdf", pdf_bytes, "application/pdf")})
        check("PDF 上传 201（状态 uploaded）", r.status_code == 201 and r.json()["status"] == "uploaded", r.text[:120])
        doc_pdf = r.json()

        txt_content = "数据结构笔记：哈希表通过散列函数将键映射到桶，冲突用链地址法解决。\n平衡二叉树的旋转操作。"
        r = http.post(f"{api}/knowledge-bases/{kb['id']}/documents", headers=auth,
                      files={"file": ("数据结构.txt", txt_content.encode("utf-8"), "text/plain")})
        check("TXT 上传 201", r.status_code == 201, r.text[:80])
        doc_txt = r.json()

        pptx_bytes = make_pptx([
            "软件工程课件第一讲：课程概述与开发流程。",
            "第二讲重点：需求分析中的用例图与边界条件识别方法。",
        ])
        r = http.post(f"{api}/knowledge-bases/{kb['id']}/documents", headers=auth,
                      files={"file": ("软件工程课件.pptx", pptx_bytes,
                                      "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
        check("PPTX 上传 201", r.status_code == 201, r.text[:80])
        doc_pptx = r.json()

        ok = wait_ready(doc_pdf["id"], auth)
        check("PDF 流水线 uploaded→parsing→indexing→ready", ok, doc_status(doc_pdf["id"], auth)["status"])

        hit = top_result(auth, "三次握手的详细过程")
        check("语义检索命中 PDF 且定位到第 2 页",
              hit is not None and hit["source_id"] == doc_pdf["id"] and hit["source_type"] == "document"
              and hit["page"] == 2, hit and f"page={hit['page']}" if hit else "no-hit")

        hit = top_result(auth, "链地址法解决冲突")
        check("语义检索命中 TXT", hit is not None and hit["source_id"] == doc_txt["id"],
              hit and f"page={hit['page']}" if hit else "no-hit")

        ok = wait_ready(doc_pptx["id"], auth)
        check("PPTX 流水线 ready", ok, doc_status(doc_pptx["id"], auth)["status"])
        hit = top_result(auth, "用例图与边界条件识别")
        check("语义检索命中 PPTX 且定位到第 2 张幻灯片",
              hit is not None and hit["source_id"] == doc_pptx["id"] and hit["page"] == 2,
              hit and f"page={hit['page']}" if hit else "no-hit")

        # --- 6. 上传校验 ---
        r = http.post(f"{api}/knowledge-bases/{kb['id']}/documents", headers=auth,
                      files={"file": ("恶意程序.exe", b"MZ\x90\x00 binary", "application/octet-stream")})
        check("扩展名白名单：.exe 422", r.status_code == 422)
        r = http.post(f"{api}/knowledge-bases/{kb['id']}/documents", headers=auth,
                      files={"file": ("伪造成.pdf", b"PK\x03\x04 not a pdf", "application/pdf")})
        check("魔数校验：非 PDF 内容 422", r.status_code == 422)
        r = http.post(f"{api}/knowledge-bases/{kb['id']}/documents", headers=auth,
                      files={"file": ("超大.txt", b"x" * (21 * 1024 * 1024), "text/plain")})
        check("大小上限 20MB：413", r.status_code == 413)
        r = http.post(f"{api}/knowledge-bases/{kb['id']}/documents", headers=auth,
                      files={"file": ("空白.txt", b"  \n\n  ", "text/plain")})
        empty_doc = r.json() if r.status_code == 201 else None
        if empty_doc:
            check("空白文档 ready（零分块不失败）", wait_ready(empty_doc["id"], auth)
                  and db["chunks"].count_documents({"document_id": empty_doc["id"]}) == 0)

        # --- 7. GridFS 去重 ---
        r = http.post(f"{api}/knowledge-bases/{kb['id']}/documents", headers=auth,
                      files={"file": ("计算机网络-副本.pdf", pdf_bytes, "application/pdf")})
        dup_doc = r.json()
        sha = db["documents"].find_one({"_id": doc_pdf["id"]})["sha256"]
        check("同字节二次上传：fs.files 单份、documents 两份",
              wait_ready(dup_doc["id"], auth)
              and db["fs.files"].count_documents({"_id": sha}) == 1
              and db["documents"].count_documents({"sha256": sha, "deleted_at": None}) == 2)

        # --- 8. reindex ---
        r = http.post(f"{api}/documents/{doc_txt['id']}/reindex", headers=auth)
        check("reindex 返回 content_version=2", r.status_code == 200 and r.json()["content_version"] == 2)
        check("reindex → 重新索引 ready(v2)", wait_ready(doc_txt["id"], auth, min_indexed_version=2))
        hit = top_result(auth, "平衡二叉树的旋转")
        check("reindex 后仍可语义检索", hit is not None and hit["source_id"] == doc_txt["id"])

        # --- 9. 删除清理 ---
        http.delete(f"{api}/notes/{note_b['id']}", headers=auth)
        ok = wait_until("笔记B向量清空", lambda: mv_rows(note_b["id"]) == []
                        and db["chunks"].count_documents({"document_id": note_b["id"]}) == 0)
        check("删笔记 → Milvus 向量与 Mongo chunks 清空", ok)

        http.delete(f"{api}/documents/{dup_doc['id']}", headers=auth)
        ok = wait_until("文档向量清空", lambda: mv_rows(dup_doc["id"]) == []
                        and db["chunks"].count_documents({"document_id": dup_doc["id"]}) == 0)
        check("删文档 → 向量与分块清空；GridFS 引用计数保留原件（同 sha 仍被引用）",
              ok and db["fs.files"].count_documents({"_id": sha}) == 1)

        # --- 10. Milvus 停机降级（§10.5）---
        mv_container = subprocess.run(
            ["docker", "ps", "--filter", "name=milvus-standalone", "--format", "{{.Names}}"],
            capture_output=True, text=True).stdout.strip().splitlines()
        if mv_container:
            subprocess.run(["docker", "stop", mv_container[0]], capture_output=True, text=True)
            try:
                r = search(auth, {"query": "页表切换", "top_k": 5})
                check("Milvus 停机 → /search 自动降级 metadata（degraded=true）",
                      r["degraded"] is True and r["mode"] == "metadata"
                      and any(x["source_type"] == "note" for x in r["results"]),
                      f"reason={r['reason'][:40]}")
            finally:
                subprocess.run(["docker", "start", mv_container[0]], capture_output=True, text=True)
                wait_until("Milvus 恢复 healthy", lambda: "healthy" in subprocess.run(
                    ["docker", "inspect", "-f", "{{.State.Health.Status}}", mv_container[0]],
                    capture_output=True, text=True).stdout, timeout=120)
                check("Milvus 恢复后语义检索正常", top_result(auth, "三次握手的详细过程") is not None)
        else:
            check("Milvus 停机降级（跳过：未找到容器）", True)

        # --- 收尾：事件队列排空 ---
        check("Redis Streams PEL 清空", wait_until(
            "pel", lambda: rds.xpending(STREAM, "graph_indexer")["pending"] == 0
            and rds.xpending(STREAM, "vector_indexer")["pending"] == 0, timeout=30))
    finally:
        stop_proc(graph)
        stop_proc(vector)
        stop_proc(publisher)
        http.close()

    print(f"\n结果：{len(PASSED)} 通过 / {len(FAILED)} 失败")
    if FAILED:
        print("失败项：", "、".join(FAILED))
        return 1
    print("阶段 4 冒烟测试全部通过 ✔")
    return 0


if __name__ == "__main__":
    sys.exit(main())
